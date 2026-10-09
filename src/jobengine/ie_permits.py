"""Companies that received Irish employment permits (flow feature 6, 9 Oct).

The Netherlands has the IND register of recognised sponsors; Ireland has no such register,
but the Department of Enterprise publishes every year how many employment permits each
company received ("Permits issued to companies", on gov.ie). A company that got permits
last year knows the process and is likely to sponsor again, so its Irish jobs rank higher in
/pending and the card says "Irish permits: N issued to <name> (DETE list)".

Config `ie_permits.url` is the gov.ie publication page (the bot follows its "companies"
download) or the file itself. The file may be an Excel sheet, a PDF or a web page: the
company column is the one with the most letters, the count the number next to it. Names
match like the IND register (legal forms and "(NL)/(IE)" stripped, register name plus
generic words, aliases). The list is read once per bot start and cached in memory.
"""

from __future__ import annotations

import io
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from urllib.parse import urljoin

from bs4 import BeautifulSoup

from jobengine.ind_register import IndRegister, parse_register_html

log = logging.getLogger("jobengine.ie_permits")

URL_KEY = "ie_permits.url"
FILE_HINT = re.compile(r"compan", re.I)
FILE_TYPES = (".xlsx", ".xls", ".pdf", ".csv")
LINE = re.compile(r"^(?P<name>.*?[A-Za-z].*?)\s+(?P<count>\d[\d,]*)\s*$")

# url -> (final url, content type, body bytes)
Getter = Callable[[str], tuple[str, str, bytes]]


@dataclass
class PermitList:
    counts: dict[str, int]  # register name -> permits issued
    register: IndRegister = field(init=False)

    def __post_init__(self) -> None:
        self.register = IndRegister.from_names(self.counts)

    def match(self, company: str | None) -> tuple[str, int] | None:
        name = self.register.match(company)
        return (name, self.counts[name]) if name else None


def _count(value: object) -> int | None:
    text = str(value or "").replace(",", "").strip()
    return int(float(text)) if re.fullmatch(r"\d+(\.0)?", text) else None


def _letters(value: object) -> int:
    return sum(ch.isalpha() for ch in str(value or ""))


def from_rows(rows: list[list[object]]) -> dict[str, int]:
    """Company -> count from table rows: the column with the most letters and the first
    numeric column after it."""
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        return {}
    width = max(len(r) for r in rows)
    letters = [sum(_letters(r[i]) for r in rows if i < len(r)) for i in range(width)]
    name_col = max(range(width), key=lambda i: letters[i])
    out: dict[str, int] = {}
    for row in rows:
        if name_col >= len(row) or _letters(row[name_col]) < 2:
            continue
        count = next((c for c in (_count(v) for v in row[name_col + 1:]) if c is not None), None)
        name = " ".join(str(row[name_col]).split())
        if count is not None and name.lower() not in ("total", "grand total"):
            out[name] = out.get(name, 0) + count
    return out


def from_xlsx(body: bytes) -> dict[str, int]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    out: dict[str, int] = {}
    for sheet in book.worksheets:
        for name, count in from_rows([list(r) for r in sheet.iter_rows(values_only=True)]).items():
            out[name] = out.get(name, 0) + count
    return out


def from_pdf(body: bytes) -> dict[str, int]:
    import pdfplumber

    out: dict[str, int] = {}
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").splitlines():
                match = LINE.match(line.strip())
                if match and match["name"].strip().lower() not in ("total", "grand total"):
                    name = " ".join(match["name"].split())
                    out[name] = out.get(name, 0) + int(match["count"].replace(",", ""))
    return out


def from_html(text: str) -> dict[str, int]:
    soup = BeautifulSoup(text, "html.parser")
    rows = [[c.get_text(" ", strip=True) for c in tr.find_all(("td", "th"))]
            for tr in soup.find_all("tr")]
    found = from_rows(rows)
    if not found:
        parse_register_html(text)  # raises ValueError with the register's own message
    return found


def download_link(page: str, base: str) -> str | None:
    """The "permits issued to companies" file linked from a gov.ie publication page."""
    soup = BeautifulSoup(page, "html.parser")
    for a in soup.find_all("a", href=True):
        href, label = str(a["href"]), a.get_text(" ", strip=True)
        if FILE_HINT.search(f"{label} {href}") and (
                href.lower().split("?")[0].endswith(FILE_TYPES) or "download" in href.lower()
                or "assets.gov.ie" in href):
            return urljoin(base, href)
    return None


def parse(url: str, kind: str, body: bytes) -> dict[str, int]:
    name = url.lower().split("?")[0]
    if "spreadsheet" in kind or "excel" in kind or name.endswith((".xlsx", ".xls")):
        return from_xlsx(body)
    if "pdf" in kind or name.endswith(".pdf"):
        return from_pdf(body)
    text = body.decode("utf-8", errors="replace")
    if "csv" in kind or name.endswith(".csv"):
        import csv

        return from_rows([list(r) for r in csv.reader(io.StringIO(text))])
    return from_html(text)


def load(url: str, get: Getter) -> PermitList:
    """Download and read the list. Raises ValueError (nothing found) or http.HttpError."""
    final, kind, body = get(url)
    if "html" in kind and not body.lstrip().startswith(b"<table"):
        link = download_link(body.decode("utf-8", errors="replace"), final)
        if link:
            final, kind, body = get(link)
    counts = parse(final, kind, body)
    if not counts:
        raise ValueError(f"no companies found in the Irish permits list ({final})")
    log.info("Irish permits list: %d companies, for example %s", len(counts),
             list(counts)[:3])
    return PermitList(counts)
